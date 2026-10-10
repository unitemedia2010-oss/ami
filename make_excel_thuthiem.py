#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Sinh file Excel CHUYÊN BIỆT: tin cho thuê căn hộ / chung cư tại Thủ Thiêm.

Tự động ghép thêm cột "Tên chủ" và "SĐT" nếu đã có file contacts.json
(tạo bằng: node contact.js --from=thu-thiem-can-ho.json --yes)

Dùng: python make_excel_thuthiem.py thu-thiem-can-ho.json Tho-thue-CH-Thu-Thiem.xlsx
"""
import json, os, re, sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "thu-thiem-can-ho.json"
DST = sys.argv[2] if len(sys.argv) > 2 else "Tho-thue-CH-Thu-Thiem.xlsx"
CONTACTS = "contacts.json"

rows = json.load(open(SRC, encoding="utf-8"))
contacts = {}
if os.path.exists(CONTACTS):
    try:
        contacts = json.load(open(CONTACTS, encoding="utf-8"))
    except Exception:
        contacts = {}

# chuẩn hoá tên dự án cho dễ đọc
def dep_hien(d):
    d = (d or "").strip()
    if not d:
        return ""
    return d.title() if d.isupper() else d


for r in rows:
    r["duAnHien"] = dep_hien(r.get("duAn"))

# sắp xếp: dự án A→Z, trong cùng dự án giá tăng dần
rows.sort(key=lambda x: (x["duAnHien"].lower(), x.get("giaTrieu", 0)))

HEAD = ["STT", "Dự án / Tòa nhà", "Giá thuê", "Diện tích (m²)", "Số PN", "Số WC",
        "Tầng", "Phường", "Địa chỉ / Đường", "Ngày đăng", "Mô tả",
        "Tên chủ", "SĐT", "Link", "Mã tin"]
WIDTH = [5, 26, 14, 11, 7, 7, 7, 20, 30, 12, 62, 20, 16, 50, 13]

H_FILL = PatternFill("solid", fgColor="0A2340")
H_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT = PatternFill("solid", fgColor="F4F8FB")

wb = Workbook()

# ---------------------------------------------------------------- sheet 1
ws = wb.active
ws.title = "Căn hộ Thủ Thiêm"
ws.append(HEAD)
for c in range(1, len(HEAD) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill, cell.font, cell.border = H_FILL, H_FONT, BORDER
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width = WIDTH[c - 1]
ws.row_dimensions[1].height = 32

so_phong = (re.compile(r"(\d+)\s*(?:pn|phòng|bed)", re.I))
dien_tich = (re.compile(r"(\d+(?:[.,]\d+)?)\s*m"))
so_wc = (re.compile(r"(\d+)\s*wc", re.I))
tang = (re.compile(r"(?:tầng|lầu)\s*(\d+)", re.I))

n = 0
for r in rows:
    n += 1
    mo = r.get("moTa") or ""
    c = contacts.get(r.get("id"), {})
    # tên chủ / sđt: thử nhiều tên trường khác nhau tuỳ phiên bản API
    ten = (c.get("tenNguoiDang") or c.get("hoTen") or c.get("tenChu") or
           c.get("nguoiDang") or c.get("name") or c.get("fullName") or "")
    sdt = (c.get("soDienThoai") or c.get("sdt") or c.get("dienThoai") or
           c.get("phone") or c.get("lienHe") or "")

    mt = dien_tich.search(mo)
    mp = so_phong.search(mo)
    mw = so_wc.search(mo)
    mg = tang.search(mo)

    ws.append([
        n, r["duAnHien"],
        f"{r.get('giaTrieu',0):,.0f} triệu/tháng".replace(",", "."),
        mt.group(1).replace(",", ".") if mt else "",
        mp.group(1) if mp else (r.get("soPhongNgu") or ""),
        mw.group(1) if mw else "",
        mg.group(1) if mg else (r.get("soTang") or ""),
        r.get("phuong", ""), r.get("duong", ""),
        (r.get("ngayDangISO") or "")[:10], mo,
        ten, sdt, r.get("link", ""), r.get("id", ""),
    ])
    rr = ws.max_row
    for cc in range(1, len(HEAD) + 1):
        ws.cell(row=rr, column=cc).border = BORDER
        if n % 2 == 0:
            ws.cell(row=rr, column=cc).fill = ALT
        ws.cell(row=rr, column=cc).alignment = Alignment(vertical="top", wrap_text=(cc in (11,)))
    if r.get("link"):
        c14 = ws.cell(row=rr, column=14)
        c14.hyperlink = r["link"]
        c14.font = Font(color="0563C1", underline="single")
    if sdt:
        ws.cell(row=rr, column=13).font = Font(bold=True, color="C00000")

ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{ws.max_row}"

# ---------------------------------------------------------------- sheet 2
ws2 = wb.create_sheet("Theo dự án")
ws2.append(["Dự án", "Số tin", "Giá thuê thấp nhất", "Giá cao nhất", "TB"])
for c in range(1, 6):
    cell = ws2.cell(row=1, column=c)
    cell.fill, cell.font, cell.border = H_FILL, H_FONT, BORDER
    cell.alignment = Alignment(horizontal="center")
for i, w in enumerate([30, 9, 18, 16, 12], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

g = {}
for r in rows:
    g.setdefault(r["duAnHien"] or "(chưa rõ)", []).append(r.get("giaTrieu", 0))
for name, prices in sorted(g.items()):
    ws2.append([name, len(prices),
                f"{min(prices):,.0f} tr".replace(",", "."),
                f"{max(prices):,.0f} tr".replace(",", "."),
                f"{sum(prices)/len(prices):,.1f} tr".replace(",", ".")])
    for c in range(1, 6):
        ws2.cell(row=ws2.max_row, column=c).border = BORDER
ws2.freeze_panes = "A2"

# ---------------------------------------------------------------- sheet 3
ws3 = wb.create_sheet("Thống kê")
ws3.column_dimensions["A"].width = 34
ws3.column_dimensions["B"].width = 16
ws3.cell(row=1, column=1, value="THỐNG KÊ THUÊ CĂN HỘ CHUNG CƯ TẠI THỦ THIÊM").font = \
    Font(bold=True, size=14, color="0A2340")
info = [
    ("Số tin", f"{len(rows)}"),
    ("Số dự án", f"{len(g)}"),
    ("Giá thấp nhất", f"{min(r.get('giaTrieu',0) for r in rows):,.0f} triệu/tháng".replace(",", ".")),
    ("Giá trung vị", f"{sorted(r.get('giaTrieu',0) for r in rows)[len(rows)//2]:,.0f} triệu/tháng".replace(",", ".")),
    ("Giá cao nhất", f"{max(r.get('giaTrieu',0) for r in rows):,.0f} triệu/tháng".replace(",", ".")),
    ("Tin đã mở khoá SĐT", f"{sum(1 for r in rows if contacts.get(r.get('id'), {}).get('ok'))}"),
]
r0 = 3
for k, v in info:
    ws3.cell(row=r0, column=1, value=k).font = Font(bold=True)
    ws3.cell(row=r0, column=1).border = BORDER
    ws3.cell(row=r0, column=2, value=v).border = BORDER
    r0 += 1

r0 += 1
ws3.cell(row=r0, column=1, value="Phân bố mức giá thuê").font = Font(bold=True, color="0A2340")
r0 += 1
thue = [r.get("giaTrieu", 0) for r in rows]
for name, lo, hi in [("Dưới 10 triệu", 0, 10), ("10 – 20 triệu", 10, 20),
                     ("20 – 30 triệu", 20, 30), ("30 – 50 triệu", 30, 50),
                     ("Trên 50 triệu", 50, 10 ** 9)]:
    c = sum(1 for v in thue if lo <= v < hi)
    ws3.cell(row=r0, column=1, value=name).border = BORDER
    ws3.cell(row=r0, column=2, value=c).border = BORDER
    r0 += 1

wb.save(DST)
have = sum(1 for r in rows if contacts.get(r.get("id"), {}).get("ok"))
print(f"OK  {DST}")
print(f"    {len(rows)} tin · {len(g)} du an · da co SĐT: {have}/{len(rows)}")